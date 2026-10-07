import pandas as pd
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from models.ticket import Ticket_Tbl


def generate_ticket_report(db):

    tickets = db.query(Ticket_Tbl).all()

    data = []

    for ticket in tickets:
        data.append({
            "Ticket ID": ticket.ticket_id,
            "Title": ticket.ticket_title,
            "Description": ticket.ticket_description,
            "Status": ticket.ticket_status,
            "User ID": ticket.ticket_user_id,
            "Started Date": ticket.started_date,
            "Updated Date": ticket.updated_date
        })

    df = pd.DataFrame(data)

    file_path = "reports/ticket_report.xlsx"

    df.to_excel(
        file_path,
        index=False,
        engine="openpyxl"
    )

    # Open the generated Excel file
    workbook = load_workbook(file_path)
    worksheet = workbook.active

    # Automatically adjust column width
    for column in worksheet.columns:
        max_length = 0
        column_letter = get_column_letter(column[0].column)

        for cell in column:
            if cell.value is not None:
                cell_length = len(str(cell.value))
                max_length = max(max_length, cell_length)

        worksheet.column_dimensions[column_letter].width = max_length + 2

    # Format date columns
    for row in range(2, worksheet.max_row + 1):
        worksheet[f"F{row}"].number_format = "yyyy-mm-dd hh:mm:ss"
        worksheet[f"G{row}"].number_format = "yyyy-mm-dd hh:mm:ss"

    workbook.save(file_path)

    return file_path